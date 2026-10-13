"""Build an A4 Word draft with heading styles, real tables, attached notes and OMML equations.
Run at project root: runtime/python/python.exe -I -B scripts/md2docx.py input.md output.docx
Pandoc is used only to convert the 17 displayed equations to native Word mathematics.
"""
import re,sys,copy,subprocess,tempfile
from pathlib import Path
from docx import Document
from docx.shared import Pt,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
src,dst=sys.argv[1],sys.argv[2]
d=Document()
for sec in d.sections:
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(2.54);sec.left_margin=sec.right_margin=Cm(2.8)
    p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
normal.element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
normal.paragraph_format.line_spacing=1.5
normal.paragraph_format.space_after=Pt(0)
for n,size in [(1,18),(2,14),(3,12)]:
    st=d.styles[f'Heading {n}'];st.font.name='Times New Roman';st.font.size=Pt(size)
    st.element.rPr.rFonts.set(qn('w:eastAsia'),'黑体');st.font.color.rgb=None
    st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(12 if n<3 else 6);st.paragraph_format.space_after=Pt(6)
# Native OMML, using the same reviewed equations as the LaTeX renderer.
formula={
1:r'\operatorname{tfidf}(t,d)=[1+\log\operatorname{tf}(t,d)]\left\{1+\log\frac{1+N}{1+\operatorname{df}(t)}\right\}',
2:r'P(y=k\mid x)=\frac{\exp(w_k\cdot x+b_k)}{\sum_j\exp(w_j\cdot x+b_j)}',
3:r'\min\;\frac12\lVert w\rVert^2+C\sum_i\max(0,1-y_i(w\cdot x_i+b))',
4:r'\operatorname{Attention}(Q,K,V)=\operatorname{softmax}\left(\frac{QK^{\mathsf T}}{\sqrt{d_k}}\right)V',
5:r'p_k=\frac{\exp(z_k)}{\sum_j\exp(z_j)},\quad k\in\{\text{消极},\text{中性},\text{积极}\}',
6:r'h=W_0x+\Delta Wx=W_0x+\frac{\alpha}{r}BAx',
7:r'P_k=\frac{\mathrm{TP}_k}{\mathrm{TP}_k+\mathrm{FP}_k},\quad R_k=\frac{\mathrm{TP}_k}{\mathrm{TP}_k+\mathrm{FN}_k},\quad \mathrm{F1}_k=\frac{2P_kR_k}{P_k+R_k}',
8:r'\text{Macro-F1}=\frac1K\sum_k\mathrm{F1}_k',
9:r'F=\frac{\sum_i n_i(\bar y_i-\bar y)^2/(g-1)}{\sum_i\sum_j(y_{ij}-\bar y_i)^2/(N-g)}',
10:r'\chi^2=\frac{(|b-c|-1)^2}{b+c}',
11:r'\tilde p_{(i)}=\max_{j\le i}\min\{1,(m-j+1)p_{(j)}\}',
12:r'\mathrm{ECE}=\sum_m\frac{|B_m|}{n}|\operatorname{acc}(B_m)-\operatorname{conf}(B_m)|',
13:r'p_k=\frac{\exp(z_k/T)}{\sum_j\exp(z_j/T)}',
14:r'\hat q=s\text{ 的第 }\lceil(n+1)(1-\alpha)\rceil\text{ 个次序统计量}',
15:r'P(y\in C(x))\ge 1-\alpha',16:r'\kappa=\frac{p_o-p_e}{1-p_e}',17:r'q=\frac{1+\sqrt{6p_o-2}}3'}
with tempfile.TemporaryDirectory() as tmp:
    p=Path(tmp)/'math.docx';command=['runtime/pandoc/pandoc.exe','-f','markdown','-t','docx','-o',str(p)]
    subprocess.run(command,input=('\n\n'.join('$$\n'+v+'\n$$' for v in formula.values())).encode(),check=True,capture_output=True)
    eq=[copy.deepcopy(x) for x in Document(p).element.body.iter(qn('m:oMathPara'))]
    assert len(eq)==17

def font(r,name='宋体',size=12,bold=None):
    r.font.name='Times New Roman';r.font.size=Pt(size)
    r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),name)
    if bold is not None:r.bold=bold

def runs(p,text,size=12,name='宋体'):
    for i,part in enumerate(re.split(r'\*\*',text)):
        if part:font(p.add_run(part),name,size,bold=(i%2==1))

def para(text,indent=True,size=12):
    p=d.add_paragraph();p.paragraph_format.widow_control=True
    if indent:p.paragraph_format.first_line_indent=Pt(size*2)
    runs(p,text,size);return p

def borders(t):
    b=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single' if edge in ['top','bottom'] else 'nil');x.set(qn('w:sz'),'8');b.append(x)
    t._tbl.tblPr.append(b)
    for c in t.rows[0].cells:
        cb=OxmlElement('w:tcBorders');x=OxmlElement('w:bottom');x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');cb.append(x);c._tc.get_or_add_tcPr().append(cb)

weights={'2-1':[.20,.23,.25,.32],'3-1':[.78,.22],'3-5':[.19,.40,.13,.28],'3-6':[.19,.65,.16],
'4-1':[.30,.07,.20,.20,.23],'4-2':[.40,.20,.20,.20],'4-3':[.20,.20,.12,.12,.12,.12,.12],
'4-4':[.20,.20,.12,.24,.24],'4-6':[.20,.20,.12,.14,.34],'4-7':[.30,.14,.14,.14,.14,.14],
'4-8':[.40,.20,.20,.20],'4-9':[.28,.10,.10,.13,.13,.13,.13],
'4-11':[.30,.10,.15,.15,.15,.15],'4-12':[.46,.54],'4-13':[.32,.12,.16,.20,.20],
'4-14':[.32,.14,.18,.18,.18],'4-15':[.40,.30,.30],'4-16':[.29,.13,.13,.15,.15,.15],'4-17':[.08,.18,.61,.13]}
lines=Path(src).read_text(encoding='utf-8').splitlines();i=0;cap='';first_heading=True
while i<len(lines):
    l=lines[i]
    if l.startswith('#'):
        level=len(l)-len(l.lstrip('#'));text=l[level:].strip();p=d.add_paragraph(text,style=f'Heading {min(level,3)}')
        if level==1:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            if not first_heading:p.paragraph_format.page_break_before=True
            first_heading=False
    elif l.startswith('> 表 '):
        cap=l[2:];p=para(cap,False,10.5);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(6)
    elif l.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            r=[v.strip() for v in lines[i].strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',v) for v in r):rows.append(r)
            i+=1
        j=i
        while j<len(lines) and not lines[j].strip():j+=1
        note=lines[j] if j<len(lines) and lines[j].startswith('注：') else None
        if note:i=j+1
        n=len(rows[0]);t=d.add_table(rows=len(rows)+(1 if note else 0),cols=n);t.autofit=False;t.alignment=WD_ALIGN_PARAGRAPH.CENTER
        num=re.search(r'表 (\d+-\d+)',cap)[1];ws=weights.get(num,[1/n]*n);ws=[v/sum(ws) for v in ws]
        for col,w in zip(t.columns,ws):col.width=Cm(15.4*w)
        for ri,row in enumerate(rows):
            pr=t.rows[ri]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
            if ri==0:pr.append(OxmlElement('w:tblHeader'))
            for ci,v in enumerate(row):
                c=t.cell(ri,ci);c.width=Cm(15.4*ws[ci]);p=c.paragraphs[0]
                p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
                p.paragraph_format.keep_with_next=(ri==0 or bool(note and ri==len(rows)-1))
                runs(p,v,10.5,'黑体' if ri==0 else '宋体')
        borders(t)
        if note:
            c=t.rows[-1].cells[0].merge(t.rows[-1].cells[-1]);p=c.paragraphs[0];runs(p,note,10)
            p.paragraph_format.line_spacing=1.15;t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        p=d.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1;p.paragraph_format.space_before=Pt(0);font(p.add_run(''),size=2)
        continue
    elif l.startswith('!['):
        m=re.match(r'!\[(.*?)\]\((.*?)\)',l);d.add_picture(m[2],width=Cm(15.0))
        p=d.paragraphs[-1];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        p=para(m[1],False,10.5);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(6)
    elif re.search(r'（2-\d+）$',l):
        k=int(re.search(r'（2-(\d+)）$',l)[1]);p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p._p.append(copy.deepcopy(eq[k-1]));font(p.add_run(f'  （2-{k}）'),size=10.5)
    elif l.startswith('```'):
        i+=1
        while not lines[i].startswith('```'):
            p=para(lines[i],False,10.5);p.paragraph_format.left_indent=Cm(.8);i+=1
    elif re.match(r'^\[\d+\]',l):
        p=para(l,False,10.5);p.paragraph_format.left_indent=Pt(21);p.paragraph_format.first_line_indent=Pt(-21);p.paragraph_format.line_spacing=1.2;p.paragraph_format.space_after=Pt(5)
    elif l.strip() and l.strip()!='---':para(l)
    i+=1
Path(dst).parent.mkdir(parents=True,exist_ok=True);d.save(dst)
print(f'Word saved: {dst}; tables={len(d.tables)}, images={len(d.inline_shapes)}, native equations={len(eq)}')
